"""P7 (T17): slide size limits and the slides-only hint in contest-mode scoring."""

import asyncio

import pytest
from fastapi import HTTPException
from pptx import Presentation

from app.config import get_settings
from app.models.schemas import SlideContent, SlideExtractionResult
from app.services import contest_scorer
from tests.test_contest_audio import TRANSCRIPT, FakeOpenAI
from tests.test_contest_scorer import QUESTION_SET
from tests.test_contest_slides import BLANK_SLIDES, SLIDES, client, fakes, post  # noqa: F401  (fixture)


def deck(pages: int, text: str = "課題", notes: str = "") -> SlideExtractionResult:
    return SlideExtractionResult(
        filename="pitch.pptx",
        slides=[SlideContent(index=i + 1, text=text, notes=notes) for i in range(pages)],
    )


def score(**kwargs):
    return asyncio.run(contest_scorer.score_materials(QUESTION_SET, **kwargs))


def score_audio(slides, tmp_path):
    audio = tmp_path / "pitch.mp3"
    audio.write_bytes(b"fake-audio-bytes")
    return asyncio.run(contest_scorer.score_audio(QUESTION_SET, str(audio), "pitch.mp3", slides=slides))


@pytest.fixture
def small_limits(monkeypatch):
    settings = get_settings()
    monkeypatch.setattr(settings, "max_contest_slide_pages", 3)
    monkeypatch.setattr(settings, "max_contest_slide_chars", 100)
    return settings


# --- (a) limits ------------------------------------------------------------------


def test_default_limits_allow_a_60_page_deck_and_reject_61(fakes):
    assert get_settings().max_contest_slide_pages == 60

    score(slides=deck(60))  # does not raise
    with pytest.raises(HTTPException) as excinfo:
        score(slides=deck(61))

    assert excinfo.value.status_code == 400
    assert "61ページ" in excinfo.value.detail
    assert "上限60ページ" in excinfo.value.detail


def test_page_limit_is_inclusive(fakes, small_limits):
    score(slides=deck(3))
    with pytest.raises(HTTPException) as excinfo:
        score(slides=deck(4))

    assert excinfo.value.status_code == 400
    assert "ページ数が多すぎます（4ページ、上限3ページ）" in excinfo.value.detail
    assert len(fakes.calls) == 1  # only the accepted deck reached Jev


def test_char_limit_counts_notes_and_is_inclusive(fakes, small_limits):
    score(slides=deck(2, text="あ" * 30, notes="い" * 20))  # 100 chars: allowed
    with pytest.raises(HTTPException) as excinfo:
        score(slides=deck(2, text="あ" * 30, notes="い" * 21))  # 102 chars

    assert excinfo.value.status_code == 400
    assert "文字数が多すぎます（102文字、上限100文字" in excinfo.value.detail
    assert "スピーカーノートを含みます" in excinfo.value.detail
    assert len(fakes.calls) == 1


def test_oversized_deck_is_rejected_before_whisper_is_called(fakes, small_limits, tmp_path):
    FakeOpenAI.calls = []
    with pytest.raises(HTTPException) as excinfo:
        score_audio(deck(4), tmp_path)

    assert excinfo.value.status_code == 400
    assert FakeOpenAI.calls == []
    assert fakes.calls == []


def test_deck_within_limits_still_scores_with_audio(fakes, small_limits, tmp_path):
    result = score_audio(deck(3), tmp_path)

    assert result.slides_included is True
    assert TRANSCRIPT in fakes.calls[0]["state"]


def test_api_oversized_pptx_is_japanese_400_without_whisper(fakes, small_limits, tmp_path):
    path = tmp_path / "big.pptx"
    presentation = Presentation()
    for i in range(4):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"スライド{i + 1}"
    presentation.save(path)
    FakeOpenAI.calls = []

    response = post([
        ("slide_file", ("big.pptx", path.read_bytes(), "application/octet-stream")),
        ("media_file", ("pitch.mp3", b"fake-audio-bytes", "audio/mpeg")),
    ])

    assert response.status_code == 400
    assert "ページ数が多すぎます（4ページ、上限3ページ）" in response.json()["detail"]
    assert FakeOpenAI.calls == []
    assert fakes.calls == []


# --- (c) slides-only hint ---------------------------------------------------------


def test_blank_transcript_with_readable_slides_suggests_slides_only(fakes):
    with pytest.raises(HTTPException) as excinfo:
        score(slides=SLIDES, transcript="   ")

    assert excinfo.value.status_code == 400
    assert "文字起こしが空" in excinfo.value.detail
    assert "スライドだけで採点することもできます" in excinfo.value.detail


@pytest.mark.parametrize("slides", [None, BLANK_SLIDES], ids=["no-slides", "slides-without-text"])
def test_blank_transcript_without_usable_slides_has_no_hint(fakes, slides):
    with pytest.raises(HTTPException) as excinfo:
        score(slides=slides, transcript="   ")

    assert "文字起こしが空" in excinfo.value.detail
    assert "スライドだけで" not in excinfo.value.detail


def test_api_silent_audio_with_slides_suggests_slides_only(fakes, tmp_path):
    from tests.test_slide_extractor import make_pptx

    path = tmp_path / "pitch.pptx"
    make_pptx(path, [("課題", "学生の8割が就活情報に困っている", "")])
    FakeOpenAI.text = ""

    response = post([
        ("slide_file", ("pitch.pptx", path.read_bytes(), "application/octet-stream")),
        ("media_file", ("silent.mp3", b"fake-audio-bytes", "audio/mpeg")),
    ])

    assert response.status_code == 400
    assert "スライドだけで採点することもできます" in response.json()["detail"]
    assert fakes.calls == []


# --- (b) regression: shapes python-pptx can't classify -----------------------------


def test_geometry_less_shape_reads_in_both_modes(fakes, tmp_path):
    from tests.test_slide_extractor import make_pptx_with_geometry_less_shape

    path = tmp_path / "nogeom.pptx"
    make_pptx_with_geometry_less_shape(path)
    deck_bytes = path.read_bytes()

    # Pitch-review tab shares the extractor.
    review = client.post("/api/slides/extract", files={"file": ("nogeom.pptx", deck_bytes, "application/octet-stream")})
    assert review.status_code == 200
    assert review.json()["slides"][0]["text"] == "形の指定が無い図形の文字"

    contest = post([("slide_file", ("nogeom.pptx", deck_bytes, "application/octet-stream"))])
    assert contest.status_code == 200
    assert "形の指定が無い図形の文字" in fakes.calls[0]["state"]
