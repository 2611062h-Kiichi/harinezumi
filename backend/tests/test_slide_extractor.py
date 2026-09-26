import pytest
from fastapi import HTTPException
from pptx import Presentation
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from app.services.slide_extractor import extract_slides


def make_pdf(path, pages_text):
    # pdfplumber (the production dependency) only reads PDFs; reportlab is a
    # test-only tool to generate one, the PDF equivalent of make_pptx() below.
    c = canvas.Canvas(str(path), pagesize=letter)
    for text in pages_text:
        if text:
            c.drawString(72, 700, text)
        c.showPage()
    c.save()


def make_pptx(path, slides):
    presentation = Presentation()
    layout = presentation.slide_layouts[1]  # title + content
    for title, body, notes in slides:
        slide = presentation.slides.add_slide(layout)
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    presentation.save(path)


def test_extracts_text_and_notes_from_pptx(tmp_path):
    path = tmp_path / "pitch.pptx"
    make_pptx(
        path,
        [
            ("課題", "学生の8割が就活情報に困っている", "ここで間を取る"),
            ("解決策", "AIが企業情報を要約する", ""),
        ],
    )

    result = extract_slides(str(path), "pitch.pptx")

    assert result.filename == "pitch.pptx"
    assert [s.index for s in result.slides] == [1, 2]
    assert "課題" in result.slides[0].text
    assert "学生の8割が就活情報に困っている" in result.slides[0].text
    assert result.slides[0].notes == "ここで間を取る"
    assert "AIが企業情報を要約する" in result.slides[1].text
    assert result.slides[1].notes == ""


def test_extracts_text_from_pdf_pages(tmp_path):
    path = tmp_path / "pitch.pdf"
    # Base14 PDF fonts don't cover Japanese without embedding a font, so this
    # uses ASCII text; Japanese extraction is already covered by the PPTX test.
    make_pdf(path, ["Problem: students struggle to find job info", "Solution: AI summarizes company data"])

    result = extract_slides(str(path), "pitch.pdf")

    assert result.filename == "pitch.pdf"
    assert [s.index for s in result.slides] == [1, 2]
    assert "Problem" in result.slides[0].text
    assert "Solution" in result.slides[1].text
    # PDFs have no speaker-notes concept, unlike PPTX.
    assert result.slides[0].notes == ""


def test_pdf_page_with_no_text_is_empty_string(tmp_path):
    path = tmp_path / "blank.pdf"
    make_pdf(path, [""])

    result = extract_slides(str(path), "blank.pdf")

    assert result.slides[0].text == ""


def test_unsupported_extension_is_rejected(tmp_path):
    path = tmp_path / "pitch.key"
    path.write_bytes(b"not a slide deck")

    with pytest.raises(HTTPException) as excinfo:
        extract_slides(str(path), "pitch.key")

    assert excinfo.value.status_code == 400


def test_extracts_text_inside_tables_and_nested_groups(tmp_path):
    from pptx.util import Inches

    path = tmp_path / "shapes.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])  # blank
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "指標"
    table.cell(0, 1).text = "数値"
    table.cell(1, 0).text = "市場規模"
    table.cell(1, 1).text = "3000億円"
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(3), Inches(2), Inches(1)).text_frame.text = "グループ内の文字"
    inner = group.shapes.add_group_shape()
    inner.shapes.add_textbox(Inches(1), Inches(4), Inches(2), Inches(1)).text_frame.text = "入れ子のグループ内の文字"
    presentation.save(path)

    text = extract_slides(str(path), "shapes.pptx").slides[0].text

    assert text.splitlines() == ["指標 | 数値", "市場規模 | 3000億円", "グループ内の文字", "入れ子のグループ内の文字"]


def test_empty_table_cells_are_skipped(tmp_path):
    from pptx.util import Inches

    path = tmp_path / "sparse.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(2, 3, Inches(1), Inches(1), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "A"
    table.cell(0, 2).text = "C"
    presentation.save(path)

    assert extract_slides(str(path), "sparse.pptx").slides[0].text == "A | C"
