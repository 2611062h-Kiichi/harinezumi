import os

import pdfplumber
from fastapi import HTTPException
from pptx import Presentation

from app.models.schemas import SlideContent, SlideExtractionResult


def extract_from_pdf(path: str, filename: str) -> SlideExtractionResult:
    slides: list[SlideContent] = []
    with pdfplumber.open(path) as pdf:
        for i, page in enumerate(pdf.pages):
            text = page.extract_text() or ""
            slides.append(SlideContent(index=i + 1, text=text.strip()))
    return SlideExtractionResult(filename=filename, slides=slides)


def extract_from_pptx(path: str, filename: str) -> SlideExtractionResult:
    presentation = Presentation(path)
    slides: list[SlideContent] = []
    for i, slide in enumerate(presentation.slides):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                texts.append(shape.text_frame.text.strip())
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        slides.append(SlideContent(index=i + 1, text="\n".join(texts), notes=notes))
    return SlideExtractionResult(filename=filename, slides=slides)


def extract_slides(path: str, filename: str) -> SlideExtractionResult:
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        return extract_from_pdf(path, filename)
    if ext == ".pptx":
        return extract_from_pptx(path, filename)
    raise HTTPException(status_code=400, detail="対応していないスライド形式です（.pdf / .pptx のみ）。")
